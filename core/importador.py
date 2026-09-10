import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from database.models import SessionLocal, Alumno, Representante, Direccion

def registrar_alumno_manual(
    cedula_alumno: str,
    nombre_alumno: str,
    apellido_alumno: str,
    fecha_nacimiento: str,
    cedula_rep: str,
    nombre_rep: str,
    apellido_rep: str,
    telefono_rep: str,
    correo_rep: str,
    anio: str = "1er Año",
    seccion: str = "A",
    estado: str = "La Guaira",
    municipio: str = "Vargas",
    parroquia: str = "Catia La Mar",
    detalle_dir: str = "Ciudad Caribia - Sector II"
) -> tuple[bool, str]:
    """Registra un alumno con su año y sección, su representante y su dirección."""
    db = SessionLocal()
    try:
        # 1. Verificar si el alumno ya existe
        alumno_existente = db.query(Alumno).filter(Alumno.cedula == cedula_alumno).first()
        if alumno_existente:
            return False, f"El alumno con cédula {cedula_alumno} ya está registrado."

        # 2. Verificar o crear representante
        rep = db.query(Representante).filter(Representante.cedula == cedula_rep).first()
        if not rep:
            # Crear dirección institucional / local
            nueva_dir = Direccion(
                estado=estado or "La Guaira",
                municipio=municipio or "Vargas",
                parroquia=parroquia or "Catia La Mar",
                detalle=detalle_dir or "Ciudad Caribia - Sector II"
            )
            db.add(nueva_dir)
            db.flush()

            rep = Representante(
                cedula=cedula_rep,
                nombre=nombre_rep,
                apellido=apellido_rep,
                telefono=telefono_rep,
                correo=correo_rep,
                id_direccion=nueva_dir.id_direccion
            )
            db.add(rep)
            db.flush()

        # 3. Crear Alumno con Año y Sección
        nuevo_alumno = Alumno(
            cedula=cedula_alumno,
            nombre=nombre_alumno,
            apellido=apellido_alumno,
            fecha_nacimiento=fecha_nacimiento,
            anio=anio or "1er Año",
            seccion=seccion or "A",
            id_representante=rep.cedula
        )
        db.add(nuevo_alumno)
        db.commit()
        return True, f"Alumno {nombre_alumno} {apellido_alumno} ({anio} - Sec. {seccion}) registrado con éxito."
    except Exception as e:
        db.rollback()
        return False, f"Error en base de datos: {str(e)}"
    finally:
        db.close()


def importar_alumnos_excel(ruta_excel: str) -> tuple[int, int, list[str]]:
    """
    Importa masivamente alumnos desde un archivo Excel (.xlsx) incluyendo Año y Sección.
    Retorna: (total_exitosos, total_fallidos, lista_de_mensajes/errores)
    """
    wb = openpyxl.load_workbook(ruta_excel)
    hoja = wb.active
    
    exitosos = 0
    fallidos = 0
    mensajes = []

    # Columnas esperadas:
    # 0: Cédula Alumno | 1: Nombre Alumno | 2: Apellido Alumno | 3: Fecha Nacimiento (DD-MM-AAAA)
    # 4: Año (ej. 3er Año) | 5: Sección (ej. A)
    # 6: Cédula Rep | 7: Nombre Rep | 8: Apellido Rep | 9: Teléfono Rep | 10: Correo Rep

    for num_fila, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        if not fila or not fila[0]: # Fila vacía
            continue

        cedula_alu = str(fila[0]).strip() if fila[0] is not None else ""
        nombre_alu = str(fila[1]).strip() if len(fila) > 1 and fila[1] is not None else ""
        apellido_alu = str(fila[2]).strip() if len(fila) > 2 and fila[2] is not None else ""
        fecha_nac = str(fila[3]).strip() if len(fila) > 3 and fila[3] is not None else ""
        anio_alu = str(fila[4]).strip() if len(fila) > 4 and fila[4] is not None else "1er Año"
        seccion_alu = str(fila[5]).strip() if len(fila) > 5 and fila[5] is not None else "A"
        
        cedula_rep = str(fila[6]).strip() if len(fila) > 6 and fila[6] is not None else ""
        nombre_rep = str(fila[7]).strip() if len(fila) > 7 and fila[7] is not None else ""
        apellido_rep = str(fila[8]).strip() if len(fila) > 8 and fila[8] is not None else ""
        tlf_rep = str(fila[9]).strip() if len(fila) > 9 and fila[9] is not None else ""
        correo_rep = str(fila[10]).strip() if len(fila) > 10 and fila[10] is not None else ""

        if not cedula_alu or not nombre_alu or not apellido_alu or not cedula_rep:
            fallidos += 1
            mensajes.append(f"Fila {num_fila}: Datos obligatorios faltantes (Cédula Alumno/Nombre/Cédula Rep).")
            continue

        ok, msg = registrar_alumno_manual(
            cedula_alumno=cedula_alu,
            nombre_alumno=nombre_alu,
            apellido_alumno=apellido_alu,
            fecha_nacimiento=fecha_nac,
            anio=anio_alu,
            seccion=seccion_alu,
            cedula_rep=cedula_rep,
            nombre_rep=nombre_rep,
            apellido_rep=apellido_rep,
            telefono_rep=tlf_rep,
            correo_rep=correo_rep
        )

        if ok:
            exitosos += 1
        else:
            fallidos += 1
            mensajes.append(f"Fila {num_fila} ({cedula_alu}): {msg}")

    return exitosos, fallidos, mensajes


def generar_plantilla_excel(ruta_guardar: str):
    """Crea la plantilla Excel estructurada con Año, Sección y datos requeridos."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matrícula SICE"

    encabezados = [
        "Cédula Alumno", 
        "Nombre Alumno", 
        "Apellido Alumno", 
        "Fecha Nacimiento (DD-MM-AAAA)",
        "Año / Nivel", 
        "Sección",
        "Cédula Representante", 
        "Nombre Representante", 
        "Apellido Representante", 
        "Teléfono Representante", 
        "Correo Representante"
    ]

    ws.append(encabezados)

    # Ejemplos diferenciando niveles (ej. 3er Año vs 5to Año)
    ejemplos = [
        [
            "V-30123456", "Juan Carlos", "Pérez", "15-03-2009",
            "3er Año", "A",
            "V-12345678", "Pedro", "Pérez", "0412-1234567", "pedro.perez@gmail.com"
        ],
        [
            "V-31987654", "Juan Carlos", "Rodríguez", "22-08-2007",
            "5to Año", "B",
            "V-13987654", "Carmen", "Rodríguez", "0424-7654321", "carmen.rod@gmail.com"
        ]
    ]

    for ej in ejemplos:
        ws.append(ej)

    # Estilos cabecera
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    wb.save(ruta_guardar)
    return True
